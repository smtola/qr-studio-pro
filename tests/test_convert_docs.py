import io
from pathlib import Path

import pytest

from src.convert.common import ConvertError, unique_path
from src.convert.documents import (csv_to_excel, excel_to_csv, excel_to_pdf, pdf_to_excel, pdf_to_word)
from src.convert.pdf_tools import images_to_pdf, merge_pdfs, parse_ranges, pdf_to_images, split_pdf


def make_xlsx(path: Path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Item", "Qty", "Price"])
    for i in range(1, 60):
        ws.append([f"Product {i}", i, i * 1.5])
    ws["A1"].font = Font(bold=True)
    ws["B1"].fill = PatternFill("solid", fgColor="FFDDDDDD")
    ws2 = wb.create_sheet("Notes")
    ws2.append(["សួស្តី", "note"])
    wb.save(path)


def make_text_pdf(path: Path, pages: int = 2):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=A4)
    for p in range(pages):
        c.setFont("Helvetica-Bold", 22)
        c.drawString(72, 780, f"Heading page {p + 1}")
        c.setFont("Helvetica", 11)
        y = 740
        for i in range(6):
            c.drawString(72, y, f"This is body line {i} of page {p + 1} with some words.")
            y -= 14
        # a ruled table
        x0, y0, cw, rh = 72, 560, 100, 22
        for r in range(4):
            for col in range(3):
                c.rect(x0 + col * cw, y0 - r * rh, cw, rh)
                c.setFont("Helvetica", 10)
                c.drawString(x0 + col * cw + 5, y0 - r * rh + 7, f"R{r}C{col}" if r else f"H{col}")
        c.showPage()
    c.save()


def test_unique_path(tmp_path):
    (tmp_path / "a.pdf").write_bytes(b"x")
    assert unique_path(tmp_path, "a", ".pdf").name == "a (2).pdf"
    assert unique_path(tmp_path, "we/ird:na*me", ".pdf").name == "we_ird_na_me.pdf"
    with pytest.raises(ConvertError):
        unique_path(tmp_path / "nope", "a", ".pdf")


def test_excel_to_pdf_basic_and_office(tmp_path):
    x = tmp_path / "book.xlsx"
    make_xlsx(x)
    for prefer in (False, True):
        out = excel_to_pdf(x, tmp_path, prefer_office=prefer)
        assert out.read_bytes().startswith(b"%PDF") and out.stat().st_size > 1000
    from pypdf import PdfReader
    assert len(PdfReader(str(unique_path(tmp_path, "book", ".pdf").with_name("book.pdf"))).pages) >= 1
    # basic renderer keeps the data
    out = excel_to_pdf(x, tmp_path, prefer_office=False)
    text = "".join(p.extract_text() for p in PdfReader(str(out)).pages)
    assert "Product 59" in text and "Sales" in text


def test_csv_roundtrip_and_pdf(tmp_path):
    c = tmp_path / "d.csv"
    c.write_text("name,qty\nأ,1\nb,2\n", encoding="utf-8")
    x = csv_to_excel(c, tmp_path)
    import openpyxl
    ws = openpyxl.load_workbook(x).active
    assert ws["B3"].value == 2
    back = excel_to_csv(x, tmp_path)
    assert back[0].read_text(encoding="utf-8-sig").splitlines()[0] == "name,qty"
    assert excel_to_pdf(c, tmp_path).suffix == ".pdf"


def test_bad_inputs(tmp_path):
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(ConvertError):
        excel_to_pdf(bad, tmp_path, prefer_office=False)
    with pytest.raises(ConvertError):
        pdf_to_word(tmp_path / "missing.pdf", tmp_path)
    junk = tmp_path / "junk.pdf"
    junk.write_bytes(b"%PDF-1.4 nonsense")
    with pytest.raises(ConvertError):
        pdf_to_word(junk, tmp_path)
    with pytest.raises(ConvertError):
        pdf_to_excel(junk, tmp_path)


def test_pdf_to_word_structure(tmp_path):
    import docx
    pdf = tmp_path / "doc.pdf"
    make_text_pdf(pdf)
    out = pdf_to_word(pdf, tmp_path)
    d = docx.Document(str(out))
    texts = [p.text for p in d.paragraphs]
    assert any("Heading page 1" in t for t in texts) and any("body line 3" in t for t in texts)
    assert any(p.style.name.startswith("Heading") for p in d.paragraphs)
    assert len(d.tables) >= 2 and d.tables[0].cell(1, 2).text == "R1C2"


def test_pdf_to_word_scanned_page_becomes_picture(tmp_path):
    from PIL import Image
    img = tmp_path / "scan.png"
    Image.new("RGB", (600, 800), "white").save(img)
    pdf = images_to_pdf([img], tmp_path, name="scan")
    notes: list[str] = []
    out = pdf_to_word(pdf, tmp_path, notes=notes)
    import docx
    assert len(docx.Document(str(out)).inline_shapes) == 1 and notes


def test_pdf_to_excel_tables_and_text(tmp_path):
    import openpyxl
    pdf = tmp_path / "t.pdf"
    make_text_pdf(pdf)
    out = pdf_to_excel(pdf, tmp_path)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames[0].startswith("p1-t1") and wb[wb.sheetnames[0]]["B2"].value == "R1C1"
    one = openpyxl.load_workbook(pdf_to_excel(pdf, tmp_path, one_sheet=True))
    assert one.sheetnames == ["Tables"]
    # a PDF with no tables falls back to text rows
    from reportlab.pdfgen import canvas
    p2 = tmp_path / "plain.pdf"
    c = canvas.Canvas(str(p2))
    c.drawString(72, 700, "just a sentence")
    c.save()
    wb2 = openpyxl.load_workbook(pdf_to_excel(p2, tmp_path))
    assert wb2.sheetnames == ["Text"] and wb2["Text"]["B2"].value == "just a sentence"


def test_images_pdf_roundtrip(tmp_path):
    from PIL import Image
    a, b = tmp_path / "a.png", tmp_path / "b.jpg"
    Image.new("RGBA", (300, 200), (255, 0, 0, 128)).save(a)
    Image.new("RGB", (200, 300), "blue").save(b)
    for page in ("fit", "a4"):
        pdf = images_to_pdf([a, b], tmp_path, name=f"imgs-{page}", page=page)
        pngs = pdf_to_images(pdf, tmp_path, dpi=72)
        assert len(pngs) == 2 and all(p.suffix == ".png" for p in pngs)
    with pytest.raises(ConvertError):
        images_to_pdf([], tmp_path)
    with pytest.raises(ConvertError):
        pdf_to_images(pdf, tmp_path, fmt="gif")


def test_merge_split(tmp_path):
    from pypdf import PdfReader
    a, b = tmp_path / "a.pdf", tmp_path / "b.pdf"
    make_text_pdf(a, 3)
    make_text_pdf(b, 2)
    m = merge_pdfs([a, b], tmp_path)
    assert len(PdfReader(str(m)).pages) == 5
    parts = split_pdf(m, tmp_path, ranges="1-2, 5")
    assert len(PdfReader(str(parts[0])).pages) == 3
    singles = split_pdf(m, tmp_path, every_page=True)
    assert len(singles) == 5
    with pytest.raises(ConvertError):
        merge_pdfs([a], tmp_path)


def test_parse_ranges():
    assert parse_ranges("1-3, 5", 6) == [0, 1, 2, 4]
    assert parse_ranges("4-", 6) == [3, 4, 5]
    assert parse_ranges("-2", 6) == [0, 1]
    for bad in ("0", "7", "3-1", "a", "", "-"):
        with pytest.raises(ConvertError):
            parse_ranges(bad, 6)
