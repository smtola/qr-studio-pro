"""Excel / Word / PDF conversions."""
from __future__ import annotations

import csv
import io
import re
import shutil
import statistics
import tempfile
from pathlib import Path
from typing import Callable

from .common import ConvertError, find_office, office_convert, require_file, unique_path

Progress = Callable[[int, int], None] | None
EXCEL_EXTS = (".xlsx", ".xlsm", ".xls", ".csv")


# ------------------------------------------------------------------------------------------ Excel -> PDF
def _rows_from_workbook(src: Path):
    """Yield (sheet_title, rows, col_widths_chars, styles) for each non-empty sheet. styles[(r,c)] = (bold, fill_hex)."""
    import openpyxl
    wb = openpyxl.load_workbook(src, data_only=True, read_only=False)
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        rows, styles = [], {}
        max_c = 0
        for r_i, row in enumerate(ws.iter_rows()):
            vals = []
            for c_i, cell in enumerate(row):
                v = cell.value
                if v is None:
                    txt = ""
                elif isinstance(v, float):
                    txt = f"{v:,.2f}".rstrip("0").rstrip(".") if abs(v) < 1e15 else str(v)
                else:
                    txt = str(v)
                vals.append(txt)
                if txt:
                    max_c = max(max_c, c_i + 1)
                    bold = bool(cell.font and cell.font.b)
                    fill = None
                    try:
                        if cell.fill and cell.fill.fill_type == "solid" and isinstance(cell.fill.fgColor.rgb, str):
                            fill = "#" + cell.fill.fgColor.rgb[-6:]
                    except Exception:  # noqa: BLE001
                        fill = None
                    if bold or fill:
                        styles[(r_i, c_i)] = (bold, fill)
            rows.append(vals)
        while rows and not any(rows[-1]):
            rows.pop()
        if not rows or max_c == 0:
            continue
        rows = [r[:max_c] + [""] * (max_c - len(r)) for r in rows]
        widths = []
        for c in range(max_c):
            letter = openpyxl.utils.get_column_letter(c + 1)
            dim = ws.column_dimensions.get(letter)
            w = dim.width if dim is not None and dim.width else None
            if not w:
                w = max((len(r[c]) for r in rows), default=8) + 2
            widths.append(min(max(float(w), 4), 60))
        yield ws.title, rows, widths, styles


def _rows_from_csv(src: Path):
    raw = src.read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeError:
            continue
    if text is None:
        raise ConvertError("The CSV encoding is not supported. Save it as UTF-8.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect)]
    if not rows:
        raise ConvertError("The CSV file is empty.")
    n = max(len(r) for r in rows)
    rows = [r + [""] * (n - len(r)) for r in rows]
    widths = [min(max(max(len(r[c]) for r in rows) + 2, 4), 60) for c in range(n)]
    yield src.stem, rows, widths, {(0, c): (True, "#EEEEEE") for c in range(n)}


def _build_pdf_fonts(sample: str):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    from ..core.fonts import has_khmer, resolve_font
    names = {}
    for bold in (False, True):
        path, _ = resolve_font("Arial", bold, has_khmer(sample))
        name = f"TLX-{'B' if bold else 'R'}-{abs(hash(path))}"
        if path:
            try:
                if name not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(name, path, subfontIndex=0))
                names[bold] = name
                continue
            except Exception:  # noqa: BLE001
                pass
        names[bold] = "Helvetica-Bold" if bold else "Helvetica"
    return names[False], names[True]


def excel_to_pdf_basic(src: Path, out: Path) -> Path:
    """Dependency-light renderer: values, bold, solid fills, column widths. Formulas show their saved results."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    sheets = list(_rows_from_csv(src) if src.suffix.lower() == ".csv" else _rows_from_workbook(src))
    if not sheets:
        raise ConvertError("The workbook has no visible cells with data.")
    sample = "".join("".join(r) for _t, rows, _w, _s in sheets for r in rows[:50])
    reg, bold = _build_pdf_fonts(sample)
    cell_style = ParagraphStyle("c", fontName=reg, fontSize=8, leading=10)
    cell_bold = ParagraphStyle("cb", parent=cell_style, fontName=bold)
    title_style = ParagraphStyle("t", fontName=bold, fontSize=12, spaceAfter=6)
    story = []
    wide = any(sum(w) > 90 for _t, _r, w, _s in sheets)
    page = landscape(A4) if wide else A4
    avail = page[0] - 36
    for si, (title, rows, widths, styles) in enumerate(sheets):
        if si:
            story.append(PageBreak())
        story += [Paragraph(escape(title), title_style), Spacer(1, 2)]
        total = sum(widths)
        scale = min(avail / (total * 4.6), 1.0) if total * 4.6 > avail else 1.0
        colw = [w * 4.6 * scale if total * 4.6 > avail else w * 4.6 for w in widths]
        if sum(colw) > avail:
            colw = [c * avail / sum(colw) for c in colw]
        data = [[Paragraph(escape(v), cell_bold if styles.get((r, c), (False, None))[0] else cell_style)
                 for c, v in enumerate(row)] for r, row in enumerate(rows)]
        ts = [("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
              ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]
        for (r, c), (_b, fill) in styles.items():
            if fill:
                ts.append(("BACKGROUND", (c, r), (c, r), colors.HexColor(fill)))
        story.append(Table(data, colWidths=colw, repeatRows=1, style=TableStyle(ts)))
    doc = SimpleDocTemplate(str(out), pagesize=page, leftMargin=18, rightMargin=18, topMargin=18, bottomMargin=18,
                            title=src.stem, creator="TL Studio Pro")
    doc.build(story)
    return out


def excel_to_pdf(src: str | Path, out_dir: str | Path, *, prefer_office: bool = True, progress: Progress = None) -> Path:
    s = require_file(src, EXCEL_EXTS)
    out = unique_path(out_dir, s.stem, ".pdf")
    if prefer_office and find_office() and s.suffix.lower() != ".csv":
        with tempfile.TemporaryDirectory(prefix="tlstudio-out-") as tmp:
            try:
                made = office_convert(s, Path(tmp), "pdf")
                shutil.move(str(made), out)
                return out
            except ConvertError:
                pass                       # fall back to the built-in renderer
    if s.suffix.lower() == ".xls":
        raise ConvertError("Old .xls files need LibreOffice. Save the file as .xlsx in Excel and try again.")
    try:
        return excel_to_pdf_basic(s, out)
    except ConvertError:
        raise
    except Exception as e:  # noqa: BLE001
        out.unlink(missing_ok=True)
        raise ConvertError(f"The spreadsheet could not be converted: {e}") from None


def word_to_pdf(src: str | Path, out_dir: str | Path, *, progress: Progress = None) -> Path:
    s = require_file(src, (".docx", ".doc", ".odt", ".rtf", ".txt"))
    out = unique_path(out_dir, s.stem, ".pdf")
    with tempfile.TemporaryDirectory(prefix="tlstudio-out-") as tmp:
        made = office_convert(s, Path(tmp), "pdf")
        shutil.move(str(made), out)
    return out


# ------------------------------------------------------------------------------------------ PDF -> Word
def _para_style(size: float, body: float) -> int | None:
    if size >= body * 1.8:
        return 1
    if size >= body * 1.4:
        return 2
    if size >= body * 1.18:
        return 3
    return None


def pdf_to_word(src: str | Path, out_dir: str | Path, *, progress: Progress = None,
                notes: list[str] | None = None) -> Path:
    """Rebuild paragraphs, headings, tables and embedded pictures in reading order. Works on PDFs with real text;
    scanned pages are inserted as images (no OCR)."""
    import docx
    import pdfplumber
    import pypdfium2 as pdfium
    from docx.shared import Inches, Pt

    s = require_file(src, (".pdf",))
    out = unique_path(out_dir, s.stem, ".docx")
    document = docx.Document()
    try:
        pdf = pdfplumber.open(str(s))
    except Exception:  # noqa: BLE001
        raise ConvertError("This PDF could not be opened (damaged or password protected).") from None
    scanned_pages = 0
    with pdf:
        n = len(pdf.pages)
        if n == 0:
            raise ConvertError("The PDF has no pages.")
        try:
            rendered = pdfium.PdfDocument(str(s))
        except Exception:  # noqa: BLE001
            rendered = None
        sizes = [w["size"] for p in pdf.pages[:20] for w in p.extract_words(extra_attrs=["size"])] or [11.0]
        body = statistics.median(sizes)
        for pi, page in enumerate(pdf.pages):
            if progress:
                progress(pi, n)
            if pi:
                document.add_page_break()
            sec = document.sections[-1]
            sec.page_width, sec.page_height = Pt(page.width), Pt(page.height)
            for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
                setattr(sec, side, Pt(36))
            tables = page.find_tables()
            tb_boxes = [t.bbox for t in tables]
            words = [w for w in page.extract_words(extra_attrs=["size", "fontname"], keep_blank_chars=False,
                                                   use_text_flow=False)
                     if not any(b[0] - 1 <= w["x0"] and w["x1"] <= b[2] + 1 and b[1] - 1 <= w["top"] and w["bottom"] <= b[3] + 1
                                for b in tb_boxes)]
            items: list[tuple[float, str, object]] = []
            # --- text lines -> paragraphs
            lines: list[list[dict]] = []
            for w in sorted(words, key=lambda x: (round(x["top"] / 3), x["x0"])):
                if lines and abs(lines[-1][0]["top"] - w["top"]) < max(w["size"] * 0.5, 2):
                    lines[-1].append(w)
                else:
                    lines.append([w])
            para: list[list[dict]] = []
            last_bottom = None

            def flush():
                nonlocal para
                if not para:
                    return
                text = " ".join(" ".join(x["text"] for x in sorted(l, key=lambda z: z["x0"])) for l in para)
                ws = [x for l in para for x in l]
                items.append((para[0][0]["top"], "p", (text, statistics.mean(x["size"] for x in ws),
                                                       sum("bold" in x["fontname"].lower() for x in ws) > len(ws) / 2,
                                                       min(x["x0"] for x in ws))))
                para = []

            for l in lines:
                top, bottom = min(x["top"] for x in l), max(x["bottom"] for x in l)
                size = statistics.mean(x["size"] for x in l)
                if last_bottom is not None and top - last_bottom > size * 0.9:
                    flush()
                if para and abs(statistics.mean(x["size"] for x in para[-1]) - size) > 1.2:
                    flush()
                para.append(l)
                last_bottom = bottom
            flush()
            # --- tables
            for t in tables:
                rows = t.extract()
                if rows and any(any(c for c in r) for r in rows):
                    items.append((t.bbox[1], "t", rows))
            # --- embedded pictures (cropped from a render of the page so position and cropping are right)
            for im in page.images:
                if rendered is None or im["width"] < 24 or im["height"] < 24:
                    continue
                try:
                    pil = rendered[pi].render(scale=150 / 72).to_pil()
                    k = 150 / 72
                    crop = pil.crop((int(im["x0"] * k), int(im["top"] * k), int(im["x1"] * k), int(im["bottom"] * k)))
                    buf = io.BytesIO()
                    crop.convert("RGB").save(buf, "PNG")
                    items.append((im["top"], "i", (buf.getvalue(), im["x1"] - im["x0"])))
                except Exception:  # noqa: BLE001
                    continue
            if not words and not tables and rendered is not None:      # scanned page: keep it as one picture
                items = []
                scanned_pages += 1
                buf = io.BytesIO()
                rendered[pi].render(scale=150 / 72).to_pil().convert("RGB").save(buf, "PNG")
                items.append((0, "i", (buf.getvalue(), page.width - 72)))
            for _top, kind, payload in sorted(items, key=lambda x: x[0]):
                if kind == "p":
                    text, size, bold, _x0 = payload
                    lvl = _para_style(size, body)
                    if lvl:
                        document.add_heading(text, level=lvl)
                    else:
                        p = document.add_paragraph()
                        r = p.add_run(text)
                        r.bold = bold
                        r.font.size = Pt(max(min(size, 28), 6))
                elif kind == "t":
                    rows = payload
                    cols = max(len(r) for r in rows)
                    tbl = document.add_table(rows=len(rows), cols=cols)
                    tbl.style = "Table Grid"
                    for ri, r in enumerate(rows):
                        for ci in range(cols):
                            tbl.cell(ri, ci).text = (r[ci] if ci < len(r) and r[ci] else "").replace("\n", " ")
                else:
                    data, w_pt = payload
                    try:
                        document.add_picture(io.BytesIO(data), width=Pt(min(w_pt, page.width - 72)))
                    except Exception:  # noqa: BLE001
                        pass
        if progress:
            progress(n, n)
    document.core_properties.title = s.stem
    document.core_properties.author = "TL Studio Pro"
    try:
        document.save(str(out))
    except OSError as e:
        raise ConvertError(f"The Word file could not be saved: {e.strerror or e}") from None
    if scanned_pages and notes is not None:
        notes.append(f"{scanned_pages} page(s) had no selectable text (scanned?) and were inserted as pictures. "
                     "Text recognition (OCR) is not included.")
    return out


# ------------------------------------------------------------------------------------------ PDF -> Excel
def _num(v: str):
    t = v.strip().replace(",", "") if isinstance(v, str) else v
    if isinstance(t, str) and re.fullmatch(r"-?\d+(\.\d+)?", t):
        return float(t) if "." in t else int(t)
    return v


def pdf_to_excel(src: str | Path, out_dir: str | Path, *, one_sheet: bool = False, progress: Progress = None) -> Path:
    """Detected tables become sheets (or stacked on one sheet). PDFs without tables are exported line by line."""
    import openpyxl
    import pdfplumber
    from openpyxl.styles import Font

    s = require_file(src, (".pdf",))
    out = unique_path(out_dir, s.stem, ".xlsx")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    found = 0
    stack = None
    try:
        pdf = pdfplumber.open(str(s))
    except Exception:  # noqa: BLE001
        raise ConvertError("This PDF could not be opened (damaged or password protected).") from None
    with pdf:
        n = len(pdf.pages)
        for pi, page in enumerate(pdf.pages):
            if progress:
                progress(pi, n)
            for ti, t in enumerate(page.find_tables()):
                rows = [[_num(c) if c is not None else None for c in r] for r in t.extract()]
                if not rows or not any(any(c not in (None, "") for c in r) for r in rows):
                    continue
                found += 1
                if one_sheet:
                    stack = stack or wb.create_sheet("Tables")
                    if stack.max_row > 1 or stack["A1"].value is not None:
                        stack.append([])
                    stack.append([f"Page {pi + 1} – table {ti + 1}"])
                    stack.cell(stack.max_row, 1).font = Font(bold=True)
                    for r in rows:
                        stack.append(r)
                else:
                    ws = wb.create_sheet(f"p{pi + 1}-t{ti + 1}"[:31])
                    for r in rows:
                        ws.append(r)
                    for c in ws[1]:
                        c.font = Font(bold=True)
        if found == 0:                                   # no tables: keep the text so nothing is lost
            ws = wb.create_sheet("Text")
            for pi, page in enumerate(pdf.pages):
                for line in (page.extract_text() or "").splitlines():
                    if line.strip():
                        ws.append([pi + 1, line])
            ws.insert_rows(1)
            ws["A1"], ws["B1"] = "Page", "Text"
            if ws.max_row <= 1:
                raise ConvertError("No text or tables were found in this PDF (it may be scanned images).")
        if progress:
            progress(n, n)
    for ws in wb.worksheets:
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 60)
    wb.save(str(out))
    return out


def excel_to_csv(src: str | Path, out_dir: str | Path, *, progress: Progress = None) -> list[Path]:
    s = require_file(src, (".xlsx", ".xlsm"))
    outs = []
    sheets = list(_rows_from_workbook(s))
    for title, rows, _w, _st in sheets:
        p = unique_path(out_dir, f"{s.stem}-{title}" if len(sheets) > 1 else s.stem, ".csv")
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f).writerows(rows)
        outs.append(p)
    if not outs:
        raise ConvertError("The workbook has no visible cells with data.")
    return outs


def csv_to_excel(src: str | Path, out_dir: str | Path, *, progress: Progress = None) -> Path:
    import openpyxl
    s = require_file(src, (".csv", ".txt"))
    out = unique_path(out_dir, s.stem, ".xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = s.stem[:31] or "Sheet1"
    for _t, rows, _w, _st in _rows_from_csv(s):
        for r in rows:
            ws.append([_num(c) for c in r])
    wb.save(str(out))
    return out
